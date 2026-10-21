from pathlib import Path
from typing import Dict
import openpyxl

from grainprint.core.database import get_session, reset_db
from grainprint.core.repository import Repository
from grainprint.core.models import MarkerAssociation


class ExcelService:

    # ========== 导出 ==========

    def export_data(self, filepath: Path | str) -> Path:
        filepath = Path(filepath)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        with get_session() as session:
            repo = Repository(session)
            self._export_crops(wb, repo)
            self._export_subspecies(wb, repo)
            self._export_varieties(wb, repo)
            self._export_markers(wb, repo)
            self._export_associations(wb, repo)

        wb.save(filepath)
        return filepath

    def _export_crops(self, wb, repo):
        ws = wb.create_sheet('粮食')
        ws.append(['ID', '名称', '英文名', '描述'])
        for c in repo.get_children('crop'):
            ws.append([c.id, c.name, c.name_en, c.description])

    def _export_subspecies(self, wb, repo):
        ws = wb.create_sheet('品类')
        ws.append(['ID', '粮食ID', '名称', '英文名', '描述'])
        for s in repo.get_children('subspecies'):
            ws.append([s.id, s.crop_id, s.name, s.name_en, s.description])

    def _export_varieties(self, wb, repo):
        ws = wb.create_sheet('品种')
        ws.append(['ID', '品类ID', '名称', '英文名'])
        for v in repo.get_children('variety'):
            ws.append([v.id, v.subspecies_id, v.name, v.name_en])

    def _export_markers(self, wb, repo):
        ws = wb.create_sheet('标志物')
        ws.append(['ID', '名称', '英文名', '结构式'])
        for m in repo.get_children('marker'):
            ws.append([m.id, m.name, m.name_en, m.structure_formula])

    def _export_associations(self, wb, repo):
        ws = wb.create_sheet('标志物关联')
        ws.append(['ID', '标志物ID', '粮食ID', '品类ID', '品种ID',
                   '来源', '筛选方法', '代谢途径', '文献来源', '备注1', '备注2'])
        with get_session() as session:
            for a in session.query(MarkerAssociation).all():
                ws.append([a.id, a.marker_id, a.crop_id, a.subspecies_id, a.variety_id,
                           a.source, a.screening_method, a.metabolic_pathway,
                           a.literature_source, a.notes, a.notes2])

    # ========== 导入 ==========

    def import_data(self, filepath: Path | str,
                    reset_database: bool = False) -> Dict[str, int]:
        filepath = Path(filepath)
        if reset_database:
            reset_db()

        wb = openpyxl.load_workbook(filepath, read_only=True)
        results = {'crops': 0, 'subspecies': 0, 'varieties': 0,
                   'markers': 0, 'associations': 0}

        with get_session() as session:
            repo = Repository(session)
            id_map = {'crops': {}, 'subspecies': {}, 'varieties': {}, 'markers': {}}

            if '粮食' in wb.sheetnames:
                results['crops'] = self._import_crops(wb['粮食'], repo, id_map)
            if '品类' in wb.sheetnames:
                results['subspecies'] = self._import_subspecies(wb['品类'], repo, id_map)
            if '品种' in wb.sheetnames:
                results['varieties'] = self._import_varieties(wb['品种'], repo, id_map)
            if '标志物' in wb.sheetnames:
                results['markers'] = self._import_markers(wb['标志物'], repo, id_map)
            if '标志物关联' in wb.sheetnames:
                results['associations'] = self._import_associations(wb['标志物关联'], repo, id_map)

        return results

    def _import_crops(self, ws, repo, id_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            old_id, name, name_en, description = (row + (None,) * 4)[:4]
            if not name:
                continue

            existing = self._find(repo, 'crop', name=name)
            if existing:
                id_map['crops'][old_id] = existing.id
            else:
                new_id = repo.create('crop', name=name, name_en=name_en,
                                     description=description)
                id_map['crops'][old_id] = new_id
                count += 1
        return count

    def _import_subspecies(self, ws, repo, id_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            old_id, old_crop_id, name, name_en, description = (row + (None,) * 5)[:5]
            if not name or not old_crop_id:
                continue

            crop_id = id_map['crops'].get(old_crop_id)
            if crop_id is None:
                print(f"⚠️  跳过品类 '{name}'：找不到对应粮食 (old_crop_id={old_crop_id})")
                continue

            existing = self._find(repo, 'subspecies', crop_id=crop_id, name=name)
            if existing:
                id_map['subspecies'][old_id] = existing.id
            else:
                new_id = repo.create('subspecies', crop_id=crop_id, name=name,
                                     name_en=name_en, description=description)
                id_map['subspecies'][old_id] = new_id
                count += 1
        return count

    def _import_varieties(self, ws, repo, id_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            old_id, old_sub_id, name, name_en = (row + (None,) * 4)[:4]
            if not name or not old_sub_id:
                continue

            sub_id = id_map['subspecies'].get(old_sub_id)
            if sub_id is None:
                print(f"⚠️  跳过品种 '{name}'：找不到对应品类 (old_sub_id={old_sub_id})")
                continue

            existing = self._find(repo, 'variety', subspecies_id=sub_id, name=name)
            if existing:
                id_map['varieties'][old_id] = existing.id
            else:
                new_id = repo.create('variety', subspecies_id=sub_id,
                                     name=name, name_en=name_en)
                id_map['varieties'][old_id] = new_id
                count += 1
        return count

    def _import_markers(self, ws, repo, id_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            old_id, name, name_en, structure_formula = (row + (None,) * 4)[:4]
            if not name:
                continue

            existing = self._find(repo, 'marker', name=name)
            if existing:
                id_map['markers'][old_id] = existing.id
            else:
                new_id = repo.create('marker', name=name, name_en=name_en,
                                     structure_formula=structure_formula)
                id_map['markers'][old_id] = new_id
                count += 1
        return count

    def _import_associations(self, ws, repo, id_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = (row + (None,) * 11)[:11]
            _, old_marker_id, old_crop_id, old_sub_id, old_var_id, \
                source, screening_method, metabolic_pathway, literature_source, notes, notes2 = row

            marker_id = id_map['markers'].get(old_marker_id)
            if marker_id is None:
                continue

            if old_crop_id:
                entity_type = 'crop'
                entity_id = id_map['crops'].get(old_crop_id)
            elif old_sub_id:
                entity_type = 'subspecies'
                entity_id = id_map['subspecies'].get(old_sub_id)
            elif old_var_id:
                entity_type = 'variety'
                entity_id = id_map['varieties'].get(old_var_id)
            else:
                continue

            if entity_id is None:
                print(f"⚠️  跳过关联：找不到 {entity_type}")
                continue

            if not self._association_exists(repo, marker_id, entity_type, entity_id):
                repo.create_marker_association(
                    marker_id=marker_id, entity_type=entity_type, entity_id=entity_id,
                    source=source, screening_method=screening_method,
                    metabolic_pathway=metabolic_pathway,
                    literature_source=literature_source, notes=notes,
                    notes2=notes2)
                count += 1
        return count

    # ========== 工具方法 ==========

    def _find(self, repo, entity_type, **filters):
        for entity in repo.get_children(entity_type):
            if all(getattr(entity, k) == v for k, v in filters.items()):
                return entity
        return None

    def _association_exists(self, repo, marker_id, entity_type, entity_id):
        associations = repo.get_marker_associations_by_entity(entity_type, entity_id)
        return any(a.marker_id == marker_id for a in associations)
